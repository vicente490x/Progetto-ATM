"""Record the ATM integration checks performed on 20 October 2026."""

from datetime import date
from pathlib import Path
from copy import copy

from openpyxl import load_workbook


book_path = Path(__file__).resolve().parents[1] / 'Consegna_esame/Casi_di_test_ATM.xlsx'
book = load_workbook(book_path)
sheet = book['Casi di test']
test_date = date(2026, 10, 20)


def result(row, scope, scenario, initial, steps, expected, observed, notes):
    values = (scope, scenario, initial, steps, expected, observed, 'Superato', test_date, notes)
    for column, value in enumerate(values, 2):
        sheet.cell(row, column).value = value
    sheet.cell(row, 9).number_format = 'dd/mm/yyyy'


sheet['A2'] = ('Collaudo locale del 20/10/2026: successi registrati solo per i casi '
               'verificati. Gli altri casi restano da eseguire.')
book['Istruzioni']['B2'] = ('31 prove pianificate; 7 superate nel collaudo locale '
                            'del 20/10/2026. Le altre restano da eseguire.')

result(
    8, 'Servizio / PostgreSQL', 'Saldo insufficiente',
    'Mario Rossi, PIN 12345; importo richiesto 999999,99 € superiore alla disponibilità.',
    'Inviare la richiesta tramite API con un nuovo UUID; controllare risposta, input e database.',
    'HTTP 409 Saldo insufficiente; nessun nuovo XML, transazione o addebito.',
    'HTTP 409: Saldo insufficiente. Input vuota; 4 transazioni e saldi totali 75000 centesimi prima e dopo.',
    'API locale /api/prelievi; UUID nuovo. La variante dal browser non è stata provata.',
)

result(
    20, 'Servizio / PostgreSQL', 'Ripetizione di richiesta già completata',
    'Un file UUID.xml già presente in processed e transazione ESEGUITA.',
    'Reinviare tramite API lo stesso ID e gli stessi dati del file processato.',
    'Risposta ESEGUITA; nessun nuovo file, transazione o addebito.',
    'Risposta ESEGUITA; 4 transazioni e saldi totali 75000 centesimi prima e dopo.',
    'Verificata idempotenza del servizio dopo il completamento. Non reinserito il file in TIBCO.',
)

result(
    21, 'Servizio / PostgreSQL', 'Identificativo riusato con importo diverso',
    'ID di una transazione già ESEGUITA; secondo importo 999999,99 €.',
    'Reinviare tramite API lo stesso ID cambiando soltanto l’importo.',
    'HTTP 409 per conflitto ID; nessun nuovo file, transazione o addebito.',
    'HTTP 409: Identificativo già usato per un’altra richiesta. Database invariato.',
    'Controllati 4 transazioni e saldi totali 75000 centesimi prima e dopo.',
)

result(
    27, 'TIBCO / PostgreSQL', 'Elaborazione di file validi',
    'BW 6.8 avviato da Business Studio; richieste UUID prenotate nel database.',
    'Inviare richieste dalla pagina e verificare processed e stato delle transazioni.',
    'XML elaborati e archiviati; transazioni ESEGUITA senza file rimasti in input.',
    'Quattro file UUID.xml in processed; quattro transazioni ESEGUITA; input vuota.',
    'Console BW del 20/10/2026. Il vecchio prelievo_valido.xml era una prova separata.',
)

result(
    29, 'TIBCO / PostgreSQL', 'File già elaborato letto di nuovo',
    'ID e53fe48c-67db-4a74-befa-d4d38c35fed7 già ESEGUITA; XML presente in processed.',
    'Riavviare BW dopo il refresh; copiare l’XML archiviato in input con lo stesso UUID.',
    'Un solo addebito; il replay esce da input e l’archivio resta coerente.',
    'Replay rimosso da input; archivio presente con SHA256 invariato. 4 transazioni ESEGUITA e saldi totali 75000 centesimi prima e dopo.',
    'Processo BW aggiornato: SpostaProcessed consente la sostituzione dello stesso UUID già archiviato.',
)

for row, test_id in ((34, 'TC_30'), (35, 'TC_31')):
    for column in range(1, 11):
        sheet.cell(row, column)._style = copy(sheet.cell(28, column)._style)
    sheet.row_dimensions[row].height = sheet.row_dimensions[28].height
    sheet.cell(row, 1).value = test_id

result(
    34, 'TIBCO / PostgreSQL', 'XML valido senza prenotazione',
    'XML conforme con nome UUID.xml; UUID assente da public.transazioni.',
    'Inserire il file in input con BW avviato; controllare cartelle, saldo e transazioni.',
    'File archiviato in error; nessuna nuova transazione o variazione del saldo.',
    'XML conforme allo XSD e spostato in error; 4 transazioni e saldi totali 75000 centesimi invariati.',
    'File fe5a73ec-15a0-404d-b82c-119931c19248.xml. Motivo JDBC non acquisito dalla Console.',
)

result(
    35, 'TIBCO / XML', 'XML non conforme: PIN mancante',
    'XML ben formato con nome UUID.xml ma privo dell’elemento PIN.',
    'Inserire il file in input con BW avviato; controllare XSD, cartelle e database.',
    'Validazione XML negativa; file in error; nessun addebito.',
    'XSD: elemento PIN mancante. File spostato in error; 4 transazioni e saldi totali 75000 centesimi invariati.',
    'File 701d7110-b8b4-4574-b3bd-68e63a925fea.xml.',
)

sheet.tables['TabellaTestATM'].ref = 'A4:J35'
sheet.print_area = 'A1:J35'
for validation in sheet.data_validations.dataValidation:
    if str(validation.sqref) == 'H5:H33':
        validation.sqref = 'H5:H35'
instructions = book['Istruzioni']
for row in range(16, 21):
    formula = instructions.cell(row, 2).value
    if isinstance(formula, str):
        instructions.cell(row, 2).value = formula.replace('A5:A33', 'A5:A35').replace('H5:H33', 'H5:H35')

book.save(book_path)
print('Registrati TC_04, TC_16, TC_17, TC_23, TC_25, TC_30 e TC_31 in', book_path)
